import base64
from unittest.mock import AsyncMock, patch

import httpx
import pytest
import click
from pathlib import Path
from click.testing import CliRunner

from src.core.api_client import OpenRouterClient
from PIL import Image
from pptx import Presentation

from src.render.cli import (
    main,
    parse_page_spec,
    collect_style_images,
    _resolve_style_paths,
)

SAMPLE_OUTLINE = """# PPT Outline: Compose CLI Test

---

## Slide 1: Cover
- **Hook:** Opening
[Visual: Title layout]

---

## Slide 2: Body
- **Point:** Detail
[Visual: Diagram]

---

## Appendix: Global Visual Requirements
- **Theme:** Navy
"""

def test_parse_page_spec_single():
    assert parse_page_spec("1") == {1}
    assert parse_page_spec("5") == {5}

def test_parse_page_spec_list():
    assert parse_page_spec("1,3,5") == {1, 3, 5}
    assert parse_page_spec(" 1 , 3 , 5 ") == {1, 3, 5}

def test_parse_page_spec_range():
    assert parse_page_spec("1-3") == {1, 2, 3}
    assert parse_page_spec("10-12") == {10, 11, 12}

def test_parse_page_spec_mixed():
    assert parse_page_spec("1,3-5,7") == {1, 3, 4, 5, 7}

def test_parse_page_spec_invalid():
    with pytest.raises(ValueError):
        parse_page_spec("0")
    with pytest.raises(ValueError):
        parse_page_spec("a")
    with pytest.raises(ValueError):
        parse_page_spec("5-1")


def test_collect_style_images_sorted(tmp_path):
    (tmp_path / "style_zebra.png").write_bytes(b"x")
    (tmp_path / "style_alpha.png").write_bytes(b"y")
    (tmp_path / "bad.gif").write_bytes(b"g")
    got = collect_style_images(tmp_path)
    assert [p.name for p in got] == ["style_alpha.png", "style_zebra.png"]


def test_resolve_style_paths_defaults_to_project_style_dir(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    style = tmp_path / "style"
    style.mkdir()
    (style / "style_cover.png").write_bytes(b"c")
    (style / "style_content.png").write_bytes(b"t")
    work = tmp_path / "work"
    work.mkdir()
    (work / "style_cover.png").write_bytes(b"old")

    got = _resolve_style_paths(Path("style"), explicit=False)
    assert got is not None
    assert [p.name for p in got] == ["style_content.png", "style_cover.png"]
    assert all(p.resolve().parent == style.resolve() for p in got)


def test_resolve_style_paths_rejects_glob(tmp_path: Path) -> None:
    with pytest.raises(click.UsageError, match="takes a directory, not a glob"):
        _resolve_style_paths(tmp_path / "*.png", explicit=True)


def test_resolve_style_paths_rejects_file(tmp_path: Path) -> None:
    plate = tmp_path / "style_cover.png"
    plate.write_bytes(b"c")
    with pytest.raises(click.UsageError, match="takes a directory, not a file"):
        _resolve_style_paths(plate, explicit=True)


def test_resolve_style_paths_explicit_missing_dir_errors(tmp_path: Path) -> None:
    with pytest.raises(click.UsageError, match="Style directory not found"):
        _resolve_style_paths(tmp_path / "nope", explicit=True)


def test_resolve_style_paths_explicit_empty_dir_errors(tmp_path: Path) -> None:
    with pytest.raises(click.UsageError, match="No style images in"):
        _resolve_style_paths(tmp_path, explicit=True)


def test_resolve_style_paths_default_missing_dir_falls_back_to_first_slide(
    tmp_path: Path,
) -> None:
    assert _resolve_style_paths(tmp_path / "nope", explicit=False) is None


def test_balance_only_prints_openrouter_credits(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, respx_mock
) -> None:
    monkeypatch.chdir(tmp_path)
    (tmp_path / ".env").write_text(
        "max_concurrent = 1\nprovider = openrouter\n\n"
        "[openrouter]\napi_key = sk-t\nimg_model = m\n",
        encoding="utf-8",
    )
    respx_mock.get(f"{OpenRouterClient.BASE_URL}/credits").mock(
        return_value=httpx.Response(
            200,
            json={"data": {"total_credits": 10.0, "total_usage": 2.0}},
        )
    )
    runner = CliRunner()
    result = runner.invoke(main, ["--balance-only"])
    assert result.exit_code == 0
    assert "OpenRouter credits" in result.output
    assert "8" in result.output


def test_balance_only_ignores_outline_flag(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, respx_mock) -> None:
    monkeypatch.chdir(tmp_path)
    (tmp_path / ".env").write_text(
        "max_concurrent = 1\nprovider = openrouter\n\n"
        "[openrouter]\napi_key = sk-t\nimg_model = m\n",
        encoding="utf-8",
    )
    outline = tmp_path / "o.md"
    outline.write_text("# x\n", encoding="utf-8")
    respx_mock.get(f"{OpenRouterClient.BASE_URL}/credits").mock(
        return_value=httpx.Response(
            200,
            json={"data": {"total_credits": 10.0, "total_usage": 2.0}},
        )
    )
    runner = CliRunner()
    result = runner.invoke(main, ["--balance-only", "--script", str(outline)])
    assert result.exit_code == 0


def test_balance_only_rejects_no_balance() -> None:
    runner = CliRunner()
    result = runner.invoke(main, ["--balance-only", "--no-balance"])
    assert result.exit_code != 0


def test_missing_outline_file_shows_error(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    (tmp_path / ".env").write_text(
        "max_concurrent = 1\nprovider = openrouter\n\n"
        "[openrouter]\napi_key = sk-t\nimg_model = m\n",
        encoding="utf-8",
    )
    runner = CliRunner()
    result = runner.invoke(
        main,
        ["--work", str(tmp_path), "--script", str(tmp_path / "missing.md")],
    )
    assert result.exit_code != 0
    assert "Script file not found" in result.output


def test_pptx_only_rebuilds_combined_pptx(tmp_path: Path) -> None:
    work = tmp_path / "work"
    work.mkdir()
    out = work / "image_test"
    out.mkdir()
    for name in ("slide_p01_v01.png", "slide_p01_v02.png", "slide_p02_v01.png"):
        Image.new("RGB", (40, 40), color="green").save(out / name)

    outline = tmp_path / "outline.md"
    outline.write_text(
        """# Deck

---

## Slide 1: Cover
[Speech: Hello slide one.]

---

## Slide 2: Body
[Speech: Hello slide two.]
""",
        encoding="utf-8",
    )

    runner = CliRunner()
    result = runner.invoke(
        main,
        [
            "--pptx-only",
            "--work",
            str(work),
            "--output",
            str(out),
            "--variant",
            "1",
            "--script",
            str(outline),
        ],
    )
    assert result.exit_code == 0, result.output
    assert "slides_test.pptx" in result.output
    assert "2 slide" in result.output
    pptx_path = work / "slides_test.pptx"
    assert pptx_path.exists()
    prs = Presentation(str(pptx_path))
    assert len(prs.slides) == 2
    assert "Hello slide one." in prs.slides[0].notes_slide.notes_text_frame.text
    assert "Hello slide two." in prs.slides[1].notes_slide.notes_text_frame.text


def test_pptx_only_requires_output(tmp_path: Path) -> None:
    runner = CliRunner()
    result = runner.invoke(main, ["--pptx-only"])
    assert result.exit_code != 0
    assert "--output" in result.output


def test_collect_style_images_empty_dir(tmp_path: Path) -> None:
    assert collect_style_images(tmp_path) == []


def test_pptx_only_conflicts_with_balance_only() -> None:
    runner = CliRunner()
    result = runner.invoke(main, ["--pptx-only", "--output", "out", "--balance-only"])
    assert result.exit_code != 0
    assert "cannot be used with --balance-only" in result.output


def test_balance_only_rejects_volcengine(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    (tmp_path / ".env").write_text(
        "max_concurrent = 1\nprovider = volcengine\n\n"
        "[volcengine]\napi_key = ark-t\nimg_model = m\n",
        encoding="utf-8",
    )
    runner = CliRunner()
    result = runner.invoke(main, ["--balance-only"])
    assert result.exit_code != 0
    assert "requires provider openrouter" in result.output


def test_compose_cli_first_slide_mode_mocked(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, mock_image_bytes
) -> None:
    monkeypatch.chdir(tmp_path)
    work = tmp_path / "work"
    work.mkdir()
    outline = work / "outline_16.md"
    outline.write_text(SAMPLE_OUTLINE, encoding="utf-8")
    (tmp_path / ".env").write_text(
        "max_concurrent = 1\nprovider = openrouter\n\n"
        "[openrouter]\napi_key = sk-t\nimg_model = m\n",
        encoding="utf-8",
    )

    out_dir = tmp_path / "slides_out"
    b64 = base64.b64encode(mock_image_bytes).decode("ascii")

    with (
        patch("src.render.cli._resolve_style_paths", return_value=None),
        patch(
            "src.render.cli.SlideImageGenerator.generate_first_slide_images",
            new=AsyncMock(return_value=[out_dir / "slide_p01_v01.png"]),
        ),
    ):
        runner = CliRunner()
        result = runner.invoke(
            main,
            [
                "--work",
                str(work),
                "--script",
                str(outline),
                "--output",
                str(out_dir),
                "--no-balance",
            ],
        )

    assert result.exit_code == 0, result.output
    assert "first slide only" in result.output
    assert "Script backup:" in result.output
    assert (out_dir / "outline_16.md").exists()
    assert (out_dir / "outline_16.md").read_text(encoding="utf-8") == SAMPLE_OUTLINE
    assert "Done. Saved 1 image(s)" in result.output


def test_compose_cli_all_slides_with_style_mocked(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, mock_image_bytes
) -> None:
    monkeypatch.chdir(tmp_path)
    work = tmp_path / "work"
    work.mkdir()
    outline = work / "outline_16.md"
    outline.write_text(SAMPLE_OUTLINE, encoding="utf-8")
    style = tmp_path / "plates"
    style.mkdir()
    (style / "style_cover.png").write_bytes(mock_image_bytes)
    (tmp_path / ".env").write_text(
        "max_concurrent = 1\nprovider = openrouter\n\n"
        "[openrouter]\napi_key = sk-t\nimg_model = m\n",
        encoding="utf-8",
    )

    out_dir = tmp_path / "slides_out"
    saved = {
        1: [out_dir / "slide_p01_v01.png"],
        2: [out_dir / "slide_p02_v01.png"],
    }

    with patch(
        "src.render.cli.SlideImageGenerator.generate_all_slide_images",
        new=AsyncMock(return_value=saved),
    ):
        runner = CliRunner()
        result = runner.invoke(
            main,
            [
                "--work",
                str(work),
                "--script",
                str(outline),
                "--style",
                str(style),
                "--output",
                str(out_dir),
                "--no-balance",
            ],
        )

    assert result.exit_code == 0, result.output
    assert "(1): style_cover.png" in result.output
    assert "Done. Saved 2 image(s)" in result.output
