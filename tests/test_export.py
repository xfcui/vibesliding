import pytest
from pathlib import Path
from PIL import Image
from pptx import Presentation

from src.core.api_client import STYLE_IMAGE_PIXEL_SIZE
from src.core.export import (
    build_contact_sheet,
    collect_slide_image_paths,
    create_pptx_from_images,
    rebuild_combined_pptx,
    save_image,
    save_style_reference_image,
    slides_by_index_from_outline,
)


def test_save_image(tmp_path):
    file_path = tmp_path / "test.png"
    data = b"fake-image-data"
    save_image(data, file_path)
    assert file_path.read_bytes() == data


def test_save_style_reference_image_downscales_to_1k(tmp_path):
    src = tmp_path / "source.png"
    Image.new("RGB", (2560, 1440), color="blue").save(src)
    out = tmp_path / "style_cover.png"
    save_style_reference_image(
        src.read_bytes(),
        out,
        target_size=STYLE_IMAGE_PIXEL_SIZE,
    )
    with Image.open(out) as img:
        assert img.size == STYLE_IMAGE_PIXEL_SIZE



def test_collect_slide_image_paths_order_and_filters(tmp_path):
    for name in (
        "slide_p02_v01.png",
        "slide_p01_v02.png",
        "slide_p01_v01.png",
        "notes.txt",
    ):
        (tmp_path / name).write_bytes(b"x")

    assert [p.name for p in collect_slide_image_paths(tmp_path)] == [
        "slide_p01_v01.png",
        "slide_p01_v02.png",
        "slide_p02_v01.png",
    ]
    assert [p.name for p in collect_slide_image_paths(tmp_path, page_filter={2})] == [
        "slide_p02_v01.png",
    ]
    assert [p.name for p in collect_slide_image_paths(tmp_path, variant_filter={1})] == [
        "slide_p01_v01.png",
        "slide_p02_v01.png",
    ]


def test_collect_slide_image_paths_empty_raises(tmp_path):
    with pytest.raises(ValueError, match="No slide images"):
        collect_slide_image_paths(tmp_path)




def test_create_pptx_from_images_includes_speaker_notes(tmp_path):
    outline = """# Deck

---

## Slide 1: Cover
- Hook
[Visual: title]
[Speech: Welcome everyone to this talk.]

---

## Slide 2: Body
- Point
[Visual: diagram]
[Speech: The key insight is simple.]
"""
    img = Image.new("RGB", (320, 180), color="red")
    img.save(tmp_path / "slide_p01_v01.png")
    img.save(tmp_path / "slide_p01_v02.png")
    img.save(tmp_path / "slide_p02_v01.png")

    pptx_path = tmp_path / "slides.pptx"
    create_pptx_from_images(
        [
            tmp_path / "slide_p01_v01.png",
            tmp_path / "slide_p01_v02.png",
            tmp_path / "slide_p02_v01.png",
        ],
        pptx_path,
        slides_by_index_from_outline(outline),
    )

    assert pptx_path.exists()
    prs = Presentation(str(pptx_path))
    assert len(prs.slides) == 3
    notes = [
        slide.notes_slide.notes_text_frame.text.strip() for slide in prs.slides
    ]
    assert notes[0] == "Welcome everyone to this talk."
    assert notes[1] == "Welcome everyone to this talk."
    assert notes[2] == "The key insight is simple."


def test_rebuild_combined_pptx(tmp_path):
    outline = """# Deck

---

## Slide 1: Cover
[Speech: Hello.]

---

## Slide 2: Body
[Speech: Next.]
"""
    image_dir = tmp_path / "image_test"
    image_dir.mkdir()
    for name in ("slide_p01_v01.png", "slide_p01_v02.png", "slide_p02_v01.png"):
        Image.new("RGB", (80, 45), color="green").save(image_dir / name)

    pptx_path, count = rebuild_combined_pptx(
        image_dir, outline, pptx_dir=tmp_path, timestamp="test", variant_filter={1}
    )
    assert count == 2
    assert pptx_path == tmp_path / "slides_test.pptx"
    assert pptx_path.exists()

    prs = Presentation(str(pptx_path))
    assert len(prs.slides) == 2
    assert prs.slides[0].notes_slide.notes_text_frame.text.strip() == "Hello."
    assert prs.slides[1].notes_slide.notes_text_frame.text.strip() == "Next."




def test_build_contact_sheet(tmp_path):
    images: list[bytes] = []
    for color in ("red", "green", "blue", "yellow"):
        buf = tmp_path / f"{color}.png"
        Image.new("RGB", (120, 80), color=color).save(buf)
        images.append(buf.read_bytes())

    sheet_path = tmp_path / "choices.png"
    build_contact_sheet(images, sheet_path, columns=2)

    assert sheet_path.exists()
    with Image.open(sheet_path) as sheet:
        assert sheet.size == (240, 160)


def test_build_contact_sheet_with_title(tmp_path):
    images: list[bytes] = []
    for color in ("red", "green"):
        buf = tmp_path / f"{color}.png"
        Image.new("RGB", (100, 60), color=color).save(buf)
        images.append(buf.read_bytes())

    sheet_path = tmp_path / "choices_titled.png"
    build_contact_sheet(images, sheet_path, columns=2, title="BASE — pick 1-2")

    assert sheet_path.exists()
    with Image.open(sheet_path) as sheet:
        assert sheet.width == 200
        assert sheet.height > 60


def test_build_contact_sheet_empty_raises(tmp_path):
    with pytest.raises(ValueError, match="At least one image"):
        build_contact_sheet([], tmp_path / "empty.png")






def test_create_pptx_chinese_notes(tmp_path):
    outline = """# PPT Outline: 测试

---

## Slide 1: 海阔天空
- 封面
[Visual: cover]
[Speech: 各位老师，各位家长，亲爱的同学们，大家下午好！]

---
"""
    img = Image.new("RGB", (320, 180), color="red")
    img.save(tmp_path / "slide_p01_v01.png")

    pptx_path = tmp_path / "slides_zh.pptx"
    create_pptx_from_images(
        [tmp_path / "slide_p01_v01.png"],
        pptx_path,
        slides_by_index_from_outline(outline),
    )

    assert pptx_path.exists()
    prs = Presentation(str(pptx_path))
    notes = prs.slides[0].notes_slide.notes_text_frame.text
    assert "各位老师" in notes
    assert "大家下午好" in notes


