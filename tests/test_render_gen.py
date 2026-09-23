from __future__ import annotations

import base64
import io

import httpx
import pytest
from pathlib import Path
from unittest.mock import AsyncMock
from PIL import Image

from src.render.gen import (
    SlideImageGenerator,
    resolve_reference_image_paths,
    resolve_style_plate_paths,
)
from src.core.api_client import OpenRouterClient
from src.core.parser import Slide, extract_global_style
from pptx import Presentation

SAMPLE_OUTLINE = """# PPT Outline: Test Deck

---

## Slide 1: Cover
- **Hook:** Opening point
[Visual: Bold title slide]
[Speech: Welcome to the deck.]

---

## Slide 2: Content
- **Point:** Main idea
[Visual: Diagram focus]
[Speech: Here is the main idea.]

---

## Appendix: Global Visual Requirements
- **Theme:** Navy #001122
"""

@pytest.fixture
def generator(client):
    return SlideImageGenerator(client)

def test_build_deck_map_lists_titles_and_neighbour_visuals(generator):
    slides = [
        Slide(index=1, title="Title 1", content="- **Label:** body\n[Visual: Dark cover; big title]"),
        Slide(index=2, title="Title 2", content="Content 2"),
        Slide(index=3, title="Title 3", content="[Visual: Split screen. Slide number: 3]"),
    ]
    context = generator._build_deck_map(slides, 2)
    assert "Slide 1: Title 1" in context
    assert "Slide 2: Title 2 (this slide)" in context
    assert "Previous, slide 1 (Title 1): Dark cover" in context
    assert "Next, slide 3 (Title 3): Split screen." in context
    assert "**Label:**" not in context
    assert "Content 2" not in context


def test_build_prompt_strips_non_content_markup(generator):
    slide = Slide(
        index=2,
        title="Plan First",
        content=(
            "<!-- write failed: 3 bullets -->\n- **Plan:** write it down\n"
            "[Style: style_content.png]\n[Visual: Split screen]\n[Speech: hello]\n\n---"
        ),
    )
    user_p, _sys_p = generator._build_prompt(slide, "Deck map")
    assert "write it down" in user_p
    for leaked in ("<!--", "[Style:", "[Speech:", "hello", "\n---"):
        assert leaked not in user_p


def test_build_prompt_role_block_and_layout_menu(generator):
    with_visual = Slide(index=2, title="T", content="- a\n[Visual: Split screen]")
    without_visual = Slide(index=2, title="T", content="- a")

    _u, sys_visual = generator._build_prompt(with_visual, "Deck map", role="transition")
    _u, sys_plain = generator._build_prompt(without_visual, "Deck map")

    assert "# THIS SLIDE" in sys_visual
    assert "Role: transition" in sys_visual
    assert "Dark curtain background" in sys_visual
    assert "# LAYOUT" not in sys_visual
    assert "# LAYOUT" in sys_plain
    assert "# THIS SLIDE" not in sys_plain


def test_build_prompt_has_on_slide_text_rule(generator):
    slide = Slide(index=1, title="T", content="[Visual: v]")
    user_p, _sys_p = generator._build_prompt(slide, "Deck map")
    assert "On-slide text" in user_p
    assert "Core insight:" in user_p

def test_extract_global_style(generator):
    slides = [
        Slide(index=1, title="Global Visual Requirements", content="Style content"),
        Slide(index=2, title="Regular Slide", content="Regular content")
    ]
    style = extract_global_style(slides)
    assert style == "Style content"

def test_build_prompt(generator):
    slide = Slide(index=1, title="Test Slide", content="Test content [Visual: A test visual]")
    user_p, sys_p = generator._build_prompt(slide, "Outline context")
    
    assert "Test Slide" in user_p
    assert "A test visual" in user_p
    assert "Outline context" in sys_p
    assert "Test content" in user_p


def test_resolve_slide_reference_images_accepts_cursor_paths(generator, tmp_path):
    ref = tmp_path / "person.png"
    ref.write_bytes(b"fake")
    slide = Slide(
        index=1,
        title="Cover",
        content=f"Intro\n[Reference: @{ref}]\n[Visual: use the photo]",
    )

    assert resolve_reference_image_paths(slide) == [ref]


def test_resolve_slide_reference_images_expands_globs(generator, tmp_path):
    (tmp_path / "peter_2.png").write_bytes(b"2")
    (tmp_path / "peter_1.png").write_bytes(b"1")
    slide = Slide(
        index=2,
        title="Hook",
        content="[Reference: peter_*.png]",
    )

    resolved = resolve_reference_image_paths(slide, outline_dir=tmp_path)

    assert [p.name for p in resolved] == ["peter_1.png", "peter_2.png"]


def test_build_prompt_injects_text_style_with_style_references(generator):
    slide = Slide(
        index=1,
        title="Cover",
        content="Cover copy\n[Visual: cinematic opener]",
    )
    global_style = "- **Theme:** Navy background\n- **Fonts:** Inter Bold 36pt"

    _user_p, sys_p = generator._build_prompt(
        slide,
        "Outline context",
        global_style=global_style,
        with_style_reference=True,
    )

    assert "VISUAL STYLE REFERENCE" in sys_p
    assert "TEXT STYLE SPEC" in sys_p
    assert "Navy background" in sys_p
    assert "Do not override the layout" in sys_p


def test_build_prompt_strips_reference_tags_and_describes_image_roles(generator, tmp_path):
    ref = tmp_path / "data.png"
    slide = Slide(
        index=1,
        title="Cover",
        content=f"Cover copy\n[Reference: {ref}]\n[Visual: use my coding figure]",
    )

    user_p, sys_p = generator._build_prompt(
        slide,
        "Outline context",
        slide_reference_paths=[ref],
        style_reference_count=2,
        with_style_reference=True,
    )

    assert "[Reference:" not in user_p
    assert "Cover copy" in user_p
    assert "Attached Slide References" in user_p
    assert "Attached image(s) 1-2: deck visual style references only" in sys_p
    assert "slide-specific photo/content references" in sys_p
    assert "Photorealistic people are allowed" in sys_p


def test_resolve_reference_image_paths_missing_raises(tmp_path: Path) -> None:
    slide = Slide(index=1, title="Cover", content="[Reference: missing.png]")
    with pytest.raises(ValueError, match="No reference image found"):
        resolve_reference_image_paths(slide, outline_dir=tmp_path)


def test_resolve_style_plate_paths_untagged_slide_is_empty(tmp_path: Path) -> None:
    slide = Slide(index=1, title="Cover", content="[Visual: opener]")
    assert resolve_style_plate_paths(slide, tmp_path) == []


def test_resolve_style_plate_paths_resolves_inside_style_dir(tmp_path: Path) -> None:
    (tmp_path / "style_content.png").write_bytes(b"c")
    (tmp_path / "style_base_content.png").write_bytes(b"b")
    slide = Slide(
        index=4,
        title="Teaching",
        content="[Style: style_content.png, style_base_content.png]\n[Visual: grid]",
    )

    resolved = resolve_style_plate_paths(slide, tmp_path)

    assert [p.name for p in resolved] == ["style_content.png", "style_base_content.png"]
    assert all(p.parent == tmp_path for p in resolved)


def test_resolve_style_plate_paths_deduplicates(tmp_path: Path) -> None:
    (tmp_path / "style_cover.png").write_bytes(b"c")
    slide = Slide(
        index=1,
        title="Cover",
        content="[Style: style_cover.png]\n[Style: style_cover.png]",
    )
    assert len(resolve_style_plate_paths(slide, tmp_path)) == 1


def test_resolve_style_plate_paths_rejects_paths(tmp_path: Path) -> None:
    nested = tmp_path / "style"
    nested.mkdir()
    (nested / "style_cover.png").write_bytes(b"c")
    slide = Slide(index=1, title="Cover", content="[Style: style/style_cover.png]")
    with pytest.raises(ValueError, match="takes a filename, not a path"):
        resolve_style_plate_paths(slide, tmp_path)


def test_resolve_style_plate_paths_missing_lists_available(tmp_path: Path) -> None:
    (tmp_path / "style_cover.png").write_bytes(b"c")
    slide = Slide(index=7, title="Teaching", content="[Style: style_nope.png]")
    with pytest.raises(ValueError, match="Available in .*style_cover.png"):
        resolve_style_plate_paths(slide, tmp_path)


def test_resolve_style_plate_paths_without_style_dir_raises(tmp_path: Path) -> None:
    slide = Slide(index=2, title="Hook", content="[Style: style_cover.png]")
    with pytest.raises(ValueError, match="no style directory is available"):
        resolve_style_plate_paths(slide, None)


def test_report_results_counts_failures(generator, capsys) -> None:
    generator._report_results(
        [
            b"ok",
            RuntimeError("boom"),
            ValueError("bad"),
            OSError("x"),
            TimeoutError("slow"),
        ],
        expected=5,
    )
    out = capsys.readouterr().out
    assert "1/5 succeeded" in out
    assert "Failure 1:" in out
    assert "... and 1 more" in out


def test_report_results_all_success(generator, capsys) -> None:
    generator._report_results([b"a", b"b"], expected=2)
    assert "Generated 2/2 image(s) successfully." in capsys.readouterr().out


def _png_bytes(color: str = "blue") -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (32, 32), color=color).save(buf, format="PNG")
    return buf.getvalue()


@pytest.mark.asyncio
async def test_generate_first_slide_images_mocked(
    generator, tmp_path: Path, respx_mock
) -> None:
    image_bytes = _png_bytes()
    b64 = base64.b64encode(image_bytes).decode("ascii")
    respx_mock.post(f"{OpenRouterClient.BASE_URL}/images").mock(
        return_value=httpx.Response(
            200,
            json={"data": [{"b64_json": b64, "media_type": "image/png"}]},
        )
    )

    out_dir = tmp_path / "image_test"
    paths = await generator.generate_first_slide_images(
        SAMPLE_OUTLINE,
        copy=2,
        output_dir=out_dir,
        work_dir=tmp_path,
        run_timestamp="test",
    )

    assert len(paths) == 2
    assert all(p.exists() for p in paths)
    prompt_path = out_dir / "slide_p01_prompt.txt"
    assert prompt_path.exists()
    prompt_text = prompt_path.read_text(encoding="utf-8")
    assert "## System prompt" in prompt_text
    assert "## User prompt" in prompt_text
    assert "Style references: none" in prompt_text
    pptx_path = tmp_path / "slides_test.pptx"
    assert pptx_path.exists()
    prs = Presentation(str(pptx_path))
    assert len(prs.slides) == 2
    assert "Welcome to the deck." in prs.slides[0].notes_slide.notes_text_frame.text


@pytest.mark.asyncio
async def test_generate_all_slide_images_mocked(
    generator, tmp_path: Path, respx_mock
) -> None:
    image_bytes = _png_bytes()
    b64 = base64.b64encode(image_bytes).decode("ascii")
    respx_mock.post(f"{OpenRouterClient.BASE_URL}/images").mock(
        return_value=httpx.Response(
            200,
            json={"data": [{"b64_json": b64, "media_type": "image/png"}]},
        )
    )

    style_ref = tmp_path / "style_cover.png"
    style_ref.write_bytes(image_bytes)
    out_dir = tmp_path / "slides"

    by_slide = await generator.generate_all_slide_images(
        SAMPLE_OUTLINE,
        style_image_paths=[style_ref],
        copy=1,
        output_dir=out_dir,
        page_filter={1, 2},
        work_dir=tmp_path,
        run_timestamp="test",
    )

    assert set(by_slide.keys()) == {1, 2}
    assert len(by_slide[1]) == 1
    assert (out_dir / "slide_p01_v01.png").exists()
    assert (out_dir / "slide_p02_v01.png").exists()
    assert (out_dir / "slide_p01_prompt.txt").exists()
    slide_two_prompt = (out_dir / "slide_p02_prompt.txt").read_text(encoding="utf-8")
    assert "style_cover.png" in slide_two_prompt
    pptx_path = tmp_path / "slides_test.pptx"
    assert pptx_path.exists()
    prs = Presentation(str(pptx_path))
    assert len(prs.slides) == 2
    assert "Here is the main idea." in prs.slides[1].notes_slide.notes_text_frame.text


@pytest.mark.asyncio
async def test_generate_all_slide_images_style_tag_overrides_role(
    generator, tmp_path: Path, respx_mock, capsys
) -> None:
    image_bytes = _png_bytes()
    b64 = base64.b64encode(image_bytes).decode("ascii")
    respx_mock.post(f"{OpenRouterClient.BASE_URL}/images").mock(
        return_value=httpx.Response(
            200,
            json={"data": [{"b64_json": b64, "media_type": "image/png"}]},
        )
    )

    style = tmp_path / "plates"
    style.mkdir()
    for name in ("style_cover.png", "style_content.png", "style_transition.png"):
        (style / name).write_bytes(image_bytes)

    # Slide 1 overrides its cover routing; slide 2 (last slide, so "ending")
    # stays on role-based routing.
    outline = SAMPLE_OUTLINE.replace(
        "[Visual: Bold title slide]",
        "[Style: style_transition.png]\n[Visual: Bold title slide]",
    )
    out_dir = tmp_path / "slides"

    await generator.generate_all_slide_images(
        outline,
        style_image_paths=sorted(style.iterdir()),
        copy=1,
        output_dir=out_dir,
        page_filter={1, 2},
        style_dir=style,
        work_dir=tmp_path,
        run_timestamp="test",
    )

    summary = capsys.readouterr().out
    assert "slide 1 (outline): style_transition.png" in summary
    assert "slide 2 (ending): style_cover.png" in summary

    slide_one_prompt = (out_dir / "slide_p01_prompt.txt").read_text(encoding="utf-8")
    assert "Style references: " in slide_one_prompt
    assert "style_transition.png" in slide_one_prompt
    assert "style_cover.png" not in slide_one_prompt


@pytest.mark.asyncio
async def test_generate_all_slide_images_unknown_style_plate_raises(
    generator, tmp_path: Path
) -> None:
    style = tmp_path / "plates"
    style.mkdir()
    (style / "style_cover.png").write_bytes(b"c")
    outline = SAMPLE_OUTLINE.replace(
        "[Visual: Bold title slide]",
        "[Style: style_missing.png]\n[Visual: Bold title slide]",
    )

    with pytest.raises(ValueError, match="Style plate not found for slide 1"):
        await generator.generate_all_slide_images(
            outline,
            style_image_paths=[style / "style_cover.png"],
            copy=1,
            output_dir=tmp_path / "out",
            style_dir=style,
        )


@pytest.mark.asyncio
async def test_generate_all_slide_images_page_filter_miss(generator, tmp_path: Path) -> None:
    style_ref = tmp_path / "style.png"
    style_ref.write_bytes(b"x")
    with pytest.raises(ValueError, match="No slides match the page filter"):
        await generator.generate_all_slide_images(
            SAMPLE_OUTLINE,
            style_image_paths=[style_ref],
            copy=1,
            output_dir=tmp_path / "out",
            page_filter={99},
        )



@pytest.mark.asyncio
async def test_generate_all_slide_images_page_filter_keeps_full_deck_roles(
    generator, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    style_dir = tmp_path / "style"
    style_dir.mkdir()
    for name in (
        "style_cover.png",
        "style_base_noncontent.png",
        "style_transition.png",
        "style_content.png",
        "style_base_content.png",
    ):
        (style_dir / name).write_bytes(b"x")
    plates = list(style_dir.iterdir())

    async def fake_parallel(prompts, on_result=None, **kwargs):
        return [b"img" for _ in prompts]

    generator.client.generate_images_parallel = AsyncMock(side_effect=fake_parallel)

    await generator.generate_all_slide_images(
        SAMPLE_OUTLINE,
        style_image_paths=plates,
        copy=1,
        output_dir=tmp_path / "out",
        page_filter={2},
        style_dir=style_dir,
    )
    logged = capsys.readouterr().out
    assert "slide 2 (ending)" in logged


def test_build_prompt_has_no_article_block(generator) -> None:
    slide = Slide(index=1, title="Stats", content="Key numbers")
    _user_p, sys_p = generator._build_prompt(
        slide,
        "Outline context",
    )
    assert "REFERENCE ARTICLES" not in sys_p
