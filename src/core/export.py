"""Export utilities - save images and assemble PPTX decks."""

from __future__ import annotations

import io
import math
import re
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.slide import Slide as PptxSlide
from pptx.util import Inches, Pt
from pptx.presentation import Presentation as PresentationType

from src.core.parser import Slide, extract_global_style, extract_speech_text, parse_markdown
from src.core.paths import (
    DEFAULT_WORK_DIR,
    slides_pptx_path,
    timestamp_from_image_dir,
    timestamp_slug,
)

SLIDE_IMAGE_PATTERN = re.compile(r"^slide_p(\d+)_v(\d+)\.png$", re.IGNORECASE)

# 16:9 widescreen (13.333" × 7.5")
PPTX_SLIDE_WIDTH = Inches(13.333333)
PPTX_SLIDE_HEIGHT = Inches(7.5)
PPTX_LATIN_FONT = "Calibri"
PPTX_EAST_ASIAN_FONT = "Microsoft YaHei"
PPTX_BODY_COLOR = RGBColor(51, 65, 85)
PPTX_NOTES_FONT_SIZE = Pt(14)


def build_contact_sheet(
    images: list[bytes],
    output_path: Path,
    *,
    columns: int = 2,
    title: str | None = None,
) -> Path:
    """Stitch candidate images into a numbered grid for visual selection.

    Args:
        images: Raw image bytes for each candidate (1..N)
        output_path: Destination PNG path
        columns: Grid columns (default 2 for a 2x2 layout when N=4)
        title: Optional header strip text (e.g. ``BASE — pick 1-4``)

    Returns:
        The saved contact sheet path

    Raises:
        ValueError: If *images* is empty or *columns* is invalid
    """
    if not images:
        raise ValueError("At least one image is required for a contact sheet")
    if columns < 1:
        raise ValueError("columns must be >= 1")

    decoded: list[Image.Image] = []
    for index, raw in enumerate(images):
        try:
            img = Image.open(io.BytesIO(raw))
            decoded.append(img.convert("RGB"))
        except Exception as exc:
            raise ValueError(f"Invalid image at index {index + 1}: {exc}") from exc

    cell_w = max(img.width for img in decoded)
    cell_h = max(img.height for img in decoded)
    rows = math.ceil(len(decoded) / columns)

    header_h = 0
    title_font = None
    if title:
        title_font = ImageFont.load_default(size=max(28, min(cell_w, cell_h) // 14))
        probe = ImageDraw.Draw(Image.new("RGB", (1, 1)))
        title_bbox = probe.textbbox((0, 0), title, font=title_font)
        header_h = (title_bbox[3] - title_bbox[1]) + 24

    grid_w = columns * cell_w
    grid_h = rows * cell_h
    sheet = Image.new("RGB", (grid_w, header_h + grid_h), color=(24, 24, 28))
    draw = ImageDraw.Draw(sheet)

    if title and title_font is not None:
        title_bbox = draw.textbbox((0, 0), title, font=title_font)
        title_w = title_bbox[2] - title_bbox[0]
        title_x = (grid_w - title_w) // 2
        title_y = 12
        draw.text((title_x, title_y), title, fill=(0, 212, 255), font=title_font)
        draw.line((0, header_h - 2, grid_w, header_h - 2), fill=(64, 64, 72), width=2)

    badge_size = max(56, min(cell_w, cell_h) // 6)
    badge_font = ImageFont.load_default(size=badge_size)
    accent = (0, 180, 220)
    border_color = (48, 48, 56)

    for index, img in enumerate(decoded):
        row, col = divmod(index, columns)
        cell_x = col * cell_w
        cell_y = header_h + row * cell_h

        draw.rectangle(
            (cell_x, cell_y, cell_x + cell_w - 1, cell_y + cell_h - 1),
            outline=border_color,
            width=2,
        )

        x = cell_x + (cell_w - img.width) // 2
        y = cell_y + (cell_h - img.height) // 2
        sheet.paste(img, (x, y))

        label = str(index + 1)
        badge_pad = badge_size // 3
        text_bbox = draw.textbbox((0, 0), label, font=badge_font)
        text_w = text_bbox[2] - text_bbox[0]
        text_h = text_bbox[3] - text_bbox[1]
        badge_w = text_w + badge_pad * 2
        badge_h = text_h + badge_pad * 2
        badge_x = cell_x + 12
        badge_y = cell_y + 12
        draw.rounded_rectangle(
            (badge_x, badge_y, badge_x + badge_w, badge_y + badge_h),
            radius=8,
            fill=accent,
        )
        draw.text(
            (badge_x + badge_pad, badge_y + badge_pad - 2),
            label,
            fill=(255, 255, 255),
            font=badge_font,
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    sheet.save(output_path, format="PNG")
    return output_path


def save_image(image_data: bytes, path: Path) -> None:
    """Save raw image bytes to file.
    
    Args:
        image_data: Raw image bytes
        path: Destination file path
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(image_data)


def save_style_reference_image(
    image_data: bytes,
    path: Path,
    *,
    target_size: tuple[int, int],
) -> None:
    """Save a style reference PNG, downscaling to *target_size* when larger."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(io.BytesIO(image_data)) as img:
        rgb = img.convert("RGB")
        if rgb.size[0] > target_size[0] or rgb.size[1] > target_size[1]:
            rgb = rgb.resize(target_size, Image.Resampling.LANCZOS)
        rgb.save(path, format="PNG")


def _slide_number_from_path(path: Path) -> int | None:
    match = SLIDE_IMAGE_PATTERN.match(path.name)
    if not match:
        return None
    return int(match.group(1))


def slides_by_index_from_outline(outline_text: str) -> dict[int, Slide]:
    """Return content slides keyed by 1-based slide index."""
    slides = parse_markdown(outline_text)
    global_style = extract_global_style(slides)
    if global_style:
        slides = [slide for slide in slides if slide.content != global_style]
    return {slide.index: slide for slide in slides}


def _slide_title_and_speech(
    slide_number: int,
    slides_by_index: dict[int, Slide],
) -> tuple[str, str]:
    slide = slides_by_index.get(slide_number)
    title = slide.title if slide is not None else f"Slide {slide_number}"
    speech = extract_speech_text(slide.content) if slide is not None else None
    return title, speech or ""


def _valid_image_paths(image_paths: list[Path]) -> list[Path]:
    valid_paths: list[Path] = []
    for path in image_paths:
        if not path.exists():
            continue
        try:
            with Image.open(path) as img:
                img.verify()
            valid_paths.append(path)
        except Exception as e:
            print(f"Warning: skipping invalid image {path}: {e}")
            continue
    return valid_paths


def _blank_slide_layout(prs: PresentationType):
    """Return the blank layout from a default python-pptx template."""
    for layout in prs.slide_layouts:
        if str(layout.name).lower() == "blank":
            return layout
    index = min(6, len(prs.slide_layouts) - 1)
    return prs.slide_layouts[index]


def _set_east_asian_typeface(run, font_name: str) -> None:
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", font_name)


def _style_notes_run(run) -> None:
    run.font.name = PPTX_LATIN_FONT
    run.font.size = PPTX_NOTES_FONT_SIZE
    run.font.color.rgb = PPTX_BODY_COLOR
    _set_east_asian_typeface(run, PPTX_EAST_ASIAN_FONT)


def _set_speaker_notes(slide: PptxSlide, text: str) -> None:
    """Write presenter speech into the slide's speaker notes."""
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    lines = text.split("\n") if text else [""]
    tf.text = lines[0]
    for line in lines[1:]:
        paragraph = tf.add_paragraph()
        paragraph.text = line
    for paragraph in tf.paragraphs:
        for run in paragraph.runs:
            _style_notes_run(run)


def create_pptx_from_images(
    image_paths: list[Path],
    output_path: Path,
    slides_by_index: dict[int, Slide] | None = None,
) -> None:
    """Combine slide images into one 16:9 PPTX with speaker notes from the outline.

    Each image becomes a full-bleed slide. ``[Speech:]`` text for that slide
    index is written into presenter notes when *slides_by_index* is provided.
    """
    valid_paths = _valid_image_paths(image_paths)
    if not valid_paths:
        raise ValueError("No valid images to process after filtering")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = PPTX_SLIDE_WIDTH
    prs.slide_height = PPTX_SLIDE_HEIGHT
    blank = _blank_slide_layout(prs)
    by_index = slides_by_index or {}

    for path in valid_paths:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(path),
            0,
            0,
            width=prs.slide_width,
            height=prs.slide_height,
        )
        slide_number = _slide_number_from_path(path)
        if slide_number is None:
            continue
        _, speech = _slide_title_and_speech(slide_number, by_index)
        _set_speaker_notes(slide, speech)

    prs.save(str(output_path))


def collect_slide_image_paths(
    output_dir: Path,
    *,
    page_filter: set[int] | None = None,
    variant_filter: set[int] | None = None,
) -> list[Path]:
    """Collect generated slide PNGs in deck order (page, then variant).

    Args:
        output_dir: Directory containing ``slide_p##_v##.png`` files
        page_filter: If set, only include these slide page numbers
        variant_filter: If set, only include these variant numbers

    Returns:
        Sorted list of matching image paths

    Raises:
        ValueError: If the directory is missing or no images match
    """
    if not output_dir.is_dir():
        raise ValueError(f"Output directory not found: {output_dir}")

    images: list[tuple[int, int, Path]] = []
    for path in output_dir.iterdir():
        if not path.is_file():
            continue
        match = SLIDE_IMAGE_PATTERN.match(path.name)
        if not match:
            continue
        page_num = int(match.group(1))
        variant_num = int(match.group(2))
        if page_filter is not None and page_num not in page_filter:
            continue
        if variant_filter is not None and variant_num not in variant_filter:
            continue
        images.append((page_num, variant_num, path))

    if not images:
        raise ValueError(
            f"No slide images found in {output_dir} "
            "(expected slide_p##_v##.png)"
        )

    images.sort(key=lambda item: (item[0], item[1]))
    return [path for _, _, path in images]


def rebuild_combined_pptx(
    image_dir: Path,
    outline_text: str | None = None,
    *,
    pptx_dir: Path | None = None,
    timestamp: str | None = None,
    page_filter: set[int] | None = None,
    variant_filter: set[int] | None = None,
) -> tuple[Path, int]:
    """Rebuild ``slides_{YYYYMMDD}.pptx`` from slide PNGs.

    Speaker notes are filled from ``[Speech:]`` tags when *outline_text* is given.
    """
    image_paths = collect_slide_image_paths(
        image_dir,
        page_filter=page_filter,
        variant_filter=variant_filter,
    )
    work = pptx_dir or DEFAULT_WORK_DIR
    ts = timestamp or timestamp_from_image_dir(image_dir) or timestamp_slug()
    pptx_path = slides_pptx_path(work, ts)
    slides_by_index = (
        slides_by_index_from_outline(outline_text) if outline_text else None
    )
    create_pptx_from_images(image_paths, pptx_path, slides_by_index)
    return pptx_path, len(image_paths)
